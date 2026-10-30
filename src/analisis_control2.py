"""
analisis_control.py  (v2)
=========================
Analisis de la respuesta de un lazo de control de temperatura (Q1 y/o Q2)
a partir de los Excel exportados (carpeta dataControl).

Que hace esta version
---------------------
1. Limpia el tiempo: el log tiene marcas repetidas y saltos (ej. 3,3,...,27,28,28,30).
   Se promedian las repetidas y se remuestrea a una malla uniforme (dt = 1 s).
   Asi las integrales (IAE, ISE, ITAE) usan un dt real y constante.
2. Separa la senal de cada planta en FASES:
      - Escalon        : desde el cambio de SetPoint hasta que aparece una perturbacion
      - Perturbacion k : BAJA (la temperatura cae) o SUBE (la temperatura sube),
                         desde su inicio hasta el inicio de la siguiente (o fin del registro)
3. Calcula las metricas de cada fase por separado (no se mezclan).
4. Exporta a Excel (metricas, tabla de perturbaciones, datos segmentados, graficas).

Uso:
    python analisis_control.py                    -> menu con los .xlsx de ./dataControl
    python analisis_control.py ruta/archivo.xlsx  -> analiza ese archivo
"""
import os
import sys
import glob
import io

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
import matplotlib.pyplot as plt
from matplotlib.patches import Patch

# ----------------------------------------------------------------------------
# PARAMETROS AJUSTABLES
# ----------------------------------------------------------------------------
CFG = dict(
    dt=1.0,                 # [s]  paso de la malla uniforme
    min_hold_sp=10.0,       # [s]  un valor de SP que dura menos se considera "glitch"
    ventana_suav=9,         # [s]  media movil usada SOLO para detectar eventos
    frac_alcanzado=0.80,    # el escalon debe llegar a este % antes de buscar perturbaciones
    umbral_min=1.0,         # [C]  desviacion minima para declarar una perturbacion
    umbral_rel=0.10,        #      ...o este % del escalon (se usa el mayor)
    ventana_evento=120.0,   # [s]  la desviacion debe ocurrir dentro de esta ventana
    t_estable=30.0,         # [s]  plateau previo exigido para aceptar una perturbacion "a favor del escalon"
    t_plano_inicio=12.0,    # [s]  PV plana este tiempo = aun no habia empezado la perturbacion
    t_asentado=60.0,        # [s]  ventana final para decidir si la respuesta al escalon ya se asento
    retiro_es_evento=False, # True: el retiro de una perturbacion (PV vuelve al nivel previo) cuenta como evento
    banda_estable=0.6,      # [C]  rango maximo (max-min) para considerar "asentada"
    t_pre=15.0,             # [s]  ventana para calcular el nivel previo a una fase
    t_ss=10.0,              # [s]  ventana para el valor final de una fase
    t_min_dentro=10.0,      # [s]  permanencia minima dentro de la banda para dar ts
)

COLOR_FASE = {"ESCALON": "#2ca02c", "BAJA": "#1f77b4", "SUBE": "#d62728"}
TXT_TIPO = {"BAJA": "BAJA (enfria)", "SUBE": "SUBE (calienta)"}

_trapz = getattr(np, "trapezoid", None) or np.trapz


# ----------------------------------------------------------------------------
# 1. LECTURA Y LIMPIEZA
# ----------------------------------------------------------------------------
def seleccionar_archivo_datos(carpeta_datos="dataControl"):
    archivos = sorted(glob.glob(os.path.join(carpeta_datos, "*.xlsx")))
    if not archivos:
        print(f"No se encontraron archivos Excel en la carpeta '{carpeta_datos}'.")
        return None
    print(f"Archivos disponibles en {carpeta_datos}:")
    for i, arch in enumerate(archivos):
        print(f"[{i}] {os.path.basename(arch)}")
    if len(archivos) == 1:
        print(f"Seleccionando unico archivo: {os.path.basename(archivos[0])}")
        return archivos[0]
    while True:
        try:
            sel = int(input(f"Seleccione el numero del archivo (0-{len(archivos) - 1}): "))
            if 0 <= sel < len(archivos):
                return archivos[sel]
            print("Seleccion invalida.")
        except ValueError:
            print("Por favor ingrese un numero valido.")


def preparar_datos(df, dt=1.0):
    """Promedia marcas de tiempo repetidas y remuestrea a malla uniforme.
    PV y corriente: interpolacion lineal.  PWM y SetPoint: retenedor (ffill)."""
    d = df.copy()
    d["Tiempo"] = pd.to_numeric(d["Tiempo"], errors="coerce")
    d = d.dropna(subset=["Tiempo"]).sort_values("Tiempo", kind="stable")
    cols = [c for c in d.columns if str(c).startswith(("Temperatura", "Corriente", "PWM", "SetPoint"))]
    g = d.groupby("Tiempo")[cols].mean()
    malla = np.arange(g.index.min(), g.index.max() + dt / 2, dt)
    out = pd.DataFrame({"Tiempo": malla})
    for c in cols:
        s = g[c].dropna()
        if s.empty:
            out[c] = np.nan
        elif str(c).startswith(("PWM", "SetPoint")):
            out[c] = s.reindex(s.index.union(malla)).ffill().reindex(malla).values
        else:
            out[c] = np.interp(malla, s.index.values, s.values)
    return out


def media_movil(x, w):
    w = max(1, int(w))
    if w == 1:
        return x.copy()
    pi = w // 2
    xp = np.pad(x, (pi, w - 1 - pi), mode="edge")
    return np.convolve(xp, np.ones(w) / w, mode="valid")


# ----------------------------------------------------------------------------
# 2. DETECCION DEL ESCALON DE SETPOINT
# ----------------------------------------------------------------------------
def detectar_escalon_sp(t, sp, cfg=CFG):
    """Devuelve idx_step (primer cambio del SP), sp_obj (valor objetivo estable),
    idx_end (siguiente cambio 'valido' de SP, o fin) y avisos sobre glitches."""
    dt = t[1] - t[0]
    n = len(sp)
    ini = np.r_[0, np.where(np.abs(np.diff(sp)) > 1e-9)[0] + 1]
    fin = np.r_[ini[1:], n]
    segs = [dict(i0=int(a), i1=int(b), valor=float(sp[a]),
                 valido=((b - a) * dt >= cfg["min_hold_sp"]) or b == n)
            for a, b in zip(ini, fin)]
    avisos = []
    if len(segs) == 1:
        avisos.append("El SetPoint no cambia en el registro: se toma el inicio como escalon.")
        return 0, float(sp[0]), n, avisos

    idx_step = segs[1]["i0"]
    validos = [s for s in segs[1:] if s["valido"]]
    sp_obj = validos[0]["valor"]
    idx_end = validos[1]["i0"] if len(validos) > 1 else n
    if len(validos) > 1:
        avisos.append(f"Hay mas cambios de SP validos; solo se analiza el primer escalon "
                      f"(hasta t={validos[1]['i0'] * dt + t[0]:.0f} s).")
    for s in segs[1:]:
        if not s["valido"]:
            avisos.append(f"SP transitorio ignorado: {s['valor']:.1f} durante "
                          f"{(s['i1'] - s['i0']) * dt:.0f} s (t={t[s['i0']]:.0f} s).")
    return idx_step, sp_obj, idx_end, avisos


# ----------------------------------------------------------------------------
# 3. DETECCION Y CLASIFICACION DE PERTURBACIONES (BAJA / SUBE)
# ----------------------------------------------------------------------------
def _inicio_movimiento(ys, i, tipo, limite, n_flat):
    """Retrocede desde el disparo `i` hasta el ultimo instante en que la PV estaba plana:
    ahi empezo realmente la desviacion."""
    sentido = 1.0 if tipo == "SUBE" else -1.0
    j, planos = i, 0
    while j > limite:
        if sentido * (ys[j] - ys[j - 1]) > 1e-6:
            planos = 0
        else:
            planos += 1
            if planos >= n_flat:
                break
        j -= 1
    return min(i, j + planos)


def _plano(w, banda):
    """True si la ventana es un plateau: poco rango Y sin tendencia (la media de la
    segunda mitad no se separa de la de la primera). Evita tomar como 'asentada'
    una senal que aun sube muy despacio."""
    if len(w) < 4:
        return False
    h = len(w) // 2
    return (np.ptp(w) <= banda) and abs(w[h:].mean() - w[:h].mean()) <= 0.5 * banda


def _asentada(ys, i, n_est, banda):
    if i - n_est < 0:
        return False
    return _plano(ys[i - n_est:i + 1], banda)


def detectar_perturbaciones(t, pv, idx_step, idx_end, y_ini, sp, cfg=CFG):
    """Busca desviaciones rapidas de la PV una vez que el escalon ya llego a
    `frac_alcanzado`.  Cada evento queda clasificado por el SIGNO de la desviacion:
        BAJA : la PV cae   (la accion de control tendra que subir)
        SUBE : la PV sube  (la accion de control tendra que bajar)

    Reglas para no confundir la respuesta natural del lazo con una perturbacion:
      - Una desviacion en el mismo sentido del escalon (ej. subir en un escalon
        ascendente) solo cuenta si antes la PV estaba asentada (plateau).
      - Una desviacion en sentido contrario se descarta si es solo la relajacion
        de un sobreimpulso hacia el SetPoint (no cruza el SP).
    """
    dt = t[1] - t[0]
    ys = media_movil(pv, cfg["ventana_suav"] / dt)
    delta = sp - y_ini
    umbral = max(cfg["umbral_min"], cfg["umbral_rel"] * abs(delta))
    if abs(delta) < 1e-9:
        return [], ys, umbral
    sgn = 1.0 if delta > 0 else -1.0

    prog = sgn * (ys - y_ini) / abs(delta)
    cand = np.where(prog[idx_step:idx_end] >= cfg["frac_alcanzado"])[0]
    if cand.size == 0:
        return [], ys, umbral
    i_a = idx_step + int(cand[0])

    W = int(cfg["ventana_evento"] / dt)
    n_est = int(cfg["t_estable"] / dt)
    n_pre = max(1, int(cfg["t_pre"] / dt))
    tol_os = max(0.5, 0.02 * abs(delta))
    n_flat = max(2, int(cfg["t_plano_inicio"] / dt))

    eventos = []
    i_busq = i_a
    i = i_a + 1
    while i < idx_end:
        j0 = max(i_busq, i - W)
        v = ys[j0:i + 1]
        caida = v.max() - ys[i]
        subida = ys[i] - v.min()

        if caida >= umbral and caida >= subida:
            tipo = "BAJA"
        elif subida >= umbral:
            tipo = "SUBE"
        else:
            i += 1
            continue
        i_on = _inicio_movimiento(ys, i, tipo, i_busq, n_flat)
        # la desviacion se mide desde donde ARRANCO el movimiento actual (no desde el
        # minimo/maximo de la ventana, que puede ser parte de la subida del escalon)
        if abs(ys[i] - ys[i_on]) < umbral:
            i += 1
            continue

        # extremo del evento: seguir mientras la PV siga en la misma direccion
        # (se detiene si la PV revierte >= umbral/2 o si se queda en plateau t_estable s)
        y_ext, i_ext, k = ys[i], i, i
        rev = 0.5 * umbral
        while k < idx_end:
            nuevo = (ys[k] < y_ext - 1e-6) if tipo == "BAJA" else (ys[k] > y_ext + 1e-6)
            if nuevo:
                y_ext, i_ext = ys[k], k
            elif abs(ys[k] - y_ext) >= rev and (ys[k] > y_ext) == (tipo == "BAJA"):
                break
            elif k - i_ext >= n_est:
                break
            k += 1

        y_pre = float(np.mean(pv[max(0, i_on - n_pre):i_on + 1]))

        a_favor = (tipo == "SUBE" and sgn > 0) or (tipo == "BAJA" and sgn < 0)
        # Evento contra el escalon que arranca del lado del sobreimpulso (PV pasada del SP):
        # es la oscilacion natural del lazo salvo que antes hubiera un plateau.
        lado_sobreimp = (sgn > 0 and y_pre > sp + tol_os) or (sgn < 0 and y_pre < sp - tol_os)
        if a_favor or lado_sobreimp:
            valido = _asentada(ys, i_on, n_est, cfg["banda_estable"])
        else:
            valido = True

        # El regreso de la PV al nivel previo de la perturbacion anterior se trata como
        # recuperacion (o retiro de la perturbacion), no como un evento nuevo.
        if valido and eventos and not cfg["retiro_es_evento"]:
            prev = eventos[-1]
            tol_rec = max(0.5 * umbral, 0.25 * abs(prev["y_pre"] - prev["y_ext"]))
            if tipo != prev["tipo"] and abs(y_ext - prev["y_pre"]) <= tol_rec:
                valido = False

        if not valido:
            i += 1
            continue

        eventos.append(dict(tipo=tipo, i_ini=int(i_on), i_ext=int(i_ext), y_pre=y_pre, y_ext=float(y_ext)))
        i_busq = int(i_ext)
        i = int(i_ext) + 1
    return eventos, ys, umbral


# ----------------------------------------------------------------------------
# 4. METRICAS POR FASE
# ----------------------------------------------------------------------------
def metricas_escalon(t, pv, pwm, corr, sp, idx_step, i_fin, y_ini, termino_por_pert, cfg=CFG):
    """Metricas SOLO de la fase de escalon [idx_step, i_fin)."""
    dt = t[1] - t[0]
    n_ss = max(1, int(cfg["t_ss"] / dt))
    n_est = max(1, int(cfg["t_asentado"] / dt))
    tt = t[idx_step:i_fin] - t[idx_step]
    y = pv[idx_step:i_fin]
    u = pwm[idx_step:i_fin]
    delta = sp - y_ini
    sgn = 1.0 if delta >= 0 else -1.0

    y_fin = float(np.mean(y[-n_ss:]))
    asentado = len(y) >= n_est and _plano(y[-n_est:], cfg["banda_estable"])
    # Valor final de referencia para tr y ts: el alcanzado (si se asento o si lo
    # corto una perturbacion); si no, el SetPoint.
    ref = y_fin if (asentado or termino_por_pert) else sp
    d_ref = ref - y_ini

    # Sobreimpulso (respecto al SetPoint y relativo al tamano del escalon)
    y_pico = y.max() if sgn > 0 else y.min()
    t_peak = float(tt[np.argmax(y) if sgn > 0 else np.argmin(y)])
    os_abs = max(0.0, sgn * (y_pico - sp))
    os_pct = os_abs / abs(delta) * 100.0 if abs(delta) > 1e-9 else 0.0

    t_rise = t_s2 = t_s5 = np.nan
    if abs(d_ref) > 1e-9:
        prog = (y - y_ini) / d_ref
        i10 = np.where(prog >= 0.10)[0]
        i90 = np.where(prog >= 0.90)[0]
        if i10.size and i90.size:
            t_rise = float(tt[i90[0]] - tt[i10[0]])

        def t_settle(p):
            fuera = np.where(np.abs(y - ref) > p * abs(d_ref))[0]
            if fuera.size == 0:
                return 0.0
            k = fuera[-1] + 1
            if k >= len(tt) or (tt[-1] - tt[k]) < cfg["t_min_dentro"]:
                return np.nan
            return float(tt[k])

        t_s2, t_s5 = t_settle(0.02), t_settle(0.05)

    e = sp - y
    iae = float(_trapz(np.abs(e), tt))
    ise = float(_trapz(e ** 2, tt))
    itae = float(_trapz(tt * np.abs(e), tt))

    i_max = i_ss = np.nan
    if corr is not None and not np.all(np.isnan(corr)):
        c = corr[idx_step:i_fin]
        i_max, i_ss = float(np.nanmax(c)), float(np.nanmean(c[-n_ss:]))

    return dict(
        t_step=float(t[idx_step]), t_fin=float(t[i_fin - 1]), duracion=float(tt[-1]),
        y_inicial=float(y_ini), sp_final=float(sp), delta_sp=float(delta),
        y_max=float(y_pico), t_peak=t_peak, sobreimpulso_abs=float(os_abs),
        sobreimpulso_pct=float(os_pct), t_rise=t_rise, t_settling_2=t_s2, t_settling_5=t_s5,
        y_ss=y_fin, error_ss=float(sp - y_fin), asentado=bool(asentado),
        termino_por="perturbacion" if termino_por_pert else "fin de registro",
        iae=iae, ise=ise, itae=itae,
        pwm_max=float(u.max()), pwm_ss=float(np.mean(u[-n_ss:])),
        tv_u=float(np.sum(np.abs(np.diff(u)))), i_max=i_max, i_ss=i_ss,
    )


def metricas_perturbacion(t, pv, pwm, ev, i0, i1, cfg=CFG):
    """Metricas de UNA perturbacion en la ventana [i0, i1). Todo se mide respecto
    al nivel previo (y_pre), de modo que el offset del lazo no contamina el resultado."""
    dt = t[1] - t[0]
    n_ss = max(1, int(cfg["t_ss"] / dt))
    n_pre = max(1, int(cfg["t_pre"] / dt))
    tipo = ev["tipo"]
    tt = t[i0:i1]
    y = pv[i0:i1]
    u = pwm[i0:i1]
    y_pre = ev["y_pre"]
    pwm_pre = float(np.mean(pwm[max(0, i0 - n_pre):i0 + 1]))

    k = int(np.argmin(y) if tipo == "BAJA" else np.argmax(y))
    y_ext, t_ext = float(y[k]), float(tt[k])
    desv = y_ext - y_pre

    ku = int(np.argmax(u) if tipo == "BAJA" else np.argmin(u))
    pwm_pico, t_pwm = float(u[ku]), float(tt[ku] - tt[0])

    y_fin = float(np.mean(y[-n_ss:]))
    rec_pct = (y_fin - y_ext) / (y_pre - y_ext) * 100.0 if abs(desv) > 1e-9 else np.nan

    def t_rec(frac):
        if abs(desv) < 1e-9:
            return np.nan
        prog = (y - y_ext) / (y_pre - y_ext)
        idx = np.where((np.arange(len(y)) > k) & (prog >= frac))[0]
        return float(tt[idx[0]] - tt[k]) if idx.size else np.nan

    d = y - y_pre
    return dict(
        tipo=tipo, t_inicio=float(tt[0]), t_fin=float(tt[-1]), duracion=float(tt[-1] - tt[0]),
        y_pre=float(y_pre), y_ext=y_ext, t_ext=t_ext, desviacion=float(desv),
        t_al_extremo=float(t_ext - tt[0]),
        pwm_pre=pwm_pre, pwm_pico=pwm_pico, delta_pwm=float(pwm_pico - pwm_pre), t_pwm_pico=t_pwm,
        y_fin=y_fin, recuperacion_pct=float(rec_pct),
        t_rec50=t_rec(0.5), t_rec90=t_rec(0.9),
        iae=float(_trapz(np.abs(d), tt)), ise=float(_trapz(d ** 2, tt)),
    )


# ----------------------------------------------------------------------------
# 5. ANALISIS COMPLETO DE UNA PLANTA
# ----------------------------------------------------------------------------
def analizar_planta(t, pv, sp, pwm, corr=None, nombre="Planta", cfg=CFG):
    idx_step, sp_obj, idx_end, avisos = detectar_escalon_sp(t, sp, cfg)
    y_ini = float(np.mean(pv[max(0, idx_step - 5):idx_step])) if idx_step > 0 else float(pv[0])

    eventos, ys, umbral = detectar_perturbaciones(t, pv, idx_step, idx_end, y_ini, sp_obj, cfg)

    i_fin_esc = eventos[0]["i_ini"] if eventos else idx_end
    esc = metricas_escalon(t, pv, pwm, corr, sp_obj, idx_step, i_fin_esc, y_ini, bool(eventos), cfg)

    perts, fases = [], [dict(nombre="Escalon", tipo="ESCALON", i0=idx_step, i1=i_fin_esc)]
    for k, ev in enumerate(eventos):
        i0 = ev["i_ini"]
        i1 = eventos[k + 1]["i_ini"] if k + 1 < len(eventos) else idx_end
        m = metricas_perturbacion(t, pv, pwm, ev, i0, i1, cfg)
        m["n"] = k + 1
        m["i_ext_abs"] = ev["i_ext"]
        perts.append(m)
        fases.append(dict(nombre=f"Pert {k + 1} {ev['tipo']}", tipo=ev["tipo"], i0=i0, i1=i1))

    etiqueta = np.full(len(t), "Pre-escalon", dtype=object)
    for f in fases:
        etiqueta[f["i0"]:f["i1"]] = f["nombre"]
    if idx_end < len(t):
        etiqueta[idx_end:] = "Fuera de analisis"

    if not eventos and not esc["asentado"]:
        avisos.append("No hay perturbaciones detectadas y la respuesta aun no se asienta "
                      "(sigue evolucionando al final del registro).")

    return dict(planta=nombre, idx_step=idx_step, sp_obj=sp_obj, y_ini=y_ini, umbral=umbral,
                escalon=esc, perturbaciones=perts, fases=fases, etiqueta=etiqueta, avisos=avisos)


# ----------------------------------------------------------------------------
# 6. SALIDA POR CONSOLA
# ----------------------------------------------------------------------------
def format_val(val, fmt=".2f", default="-"):
    if val is None:
        return default
    try:
        if np.isnan(val):
            return default
    except TypeError:
        return str(val)
    return f"{val:{fmt}}"


def limpiar_numero(x, dec=2):
    if x is None:
        return None
    if isinstance(x, (bool, np.bool_)):
        return "Si" if x else "No"
    if isinstance(x, str):
        return x
    try:
        if np.isnan(x):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(x, (int, np.integer)):
        return int(x)
    if isinstance(x, (float, np.floating)):
        r = round(float(x), dec)
        return int(r) if r.is_integer() else r
    return x


FILAS_ESCALON = [
    ("SetPoint [C]", "sp_final", 2),
    ("Temp. inicial [C]", "y_inicial", 2),
    ("Delta SP [C]", "delta_sp", 2),
    ("Inicio escalon [s]", "t_step", 1),
    ("Fin fase escalon [s]", "t_fin", 1),
    ("Fase termino por", "termino_por", 0),
    ("Respuesta asentada", "asentado", 0),
    ("Temp. maxima [C]", "y_max", 2),
    ("Tiempo pico [s]", "t_peak", 1),
    ("Sobreimpulso [C]", "sobreimpulso_abs", 2),
    ("Sobreimpulso [% del escalon]", "sobreimpulso_pct", 2),
    ("Tiempo subida tr 10-90% [s]", "t_rise", 1),
    ("Tiempo establ. 5% [s]", "t_settling_5", 1),
    ("Tiempo establ. 2% [s]", "t_settling_2", 1),
    ("Temp. final fase Yss [C]", "y_ss", 2),
    ("Error vs SP al final [C]", "error_ss", 2),
    ("IAE", "iae", 2),
    ("ISE", "ise", 2),
    ("ITAE", "itae", 1),
    ("PWM maximo [%]", "pwm_max", 2),
    ("PWM final fase [%]", "pwm_ss", 2),
    ("Variacion total TV(u)", "tv_u", 2),
    ("Corriente max [mA]", "i_max", 1),
    ("Corriente final fase [mA]", "i_ss", 1),
]

FILAS_PERT = [
    ("N", "n", 0), ("Tipo", "tipo", 0), ("Inicio [s]", "t_inicio", 1), ("Duracion ventana [s]", "duracion", 1),
    ("Temp. previa [C]", "y_pre", 2), ("Temp. extremo [C]", "y_ext", 2),
    ("Desviacion [C]", "desviacion", 2), ("Tiempo al extremo [s]", "t_al_extremo", 1),
    ("PWM previo [%]", "pwm_pre", 2), ("PWM reaccion [%]", "pwm_pico", 2),
    ("Delta PWM [%]", "delta_pwm", 2), ("t al PWM pico [s]", "t_pwm_pico", 1),
    ("Recuperado al final [%]", "recuperacion_pct", 1),
    ("t recup. 50% [s]", "t_rec50", 1), ("t recup. 90% [s]", "t_rec90", 1),
    ("IAE desviacion", "iae", 2), ("ISE desviacion", "ise", 2),
]


def _fmt_celda(v, dec):
    if isinstance(v, (bool, np.bool_)):
        return "Si" if v else "No"
    if isinstance(v, str):
        return TXT_TIPO.get(v, v)
    return format_val(v, f".{dec}f")


def imprimir_resumen_control(resultados, df_params=None):
    print("\n" + "=" * 70)
    print("                     RESUMEN DEL CONTROL")
    print("=" * 70)
    if df_params is not None and not df_params.empty:
        print("\n--- Parametros PID (configurados) ---")
        print(df_params.to_string(index=False))

    nombres = [r["planta"] for r in resultados]
    print("\n--- FASE 1: Respuesta al escalon de SetPoint ---")
    hdr = f"{'Metrica':<32}" + "".join(f" | {n:<16}" for n in nombres)
    print(hdr)
    print("-" * len(hdr))
    for label, key, dec in FILAS_ESCALON:
        print(f"{label:<32}" + "".join(f" | {_fmt_celda(r['escalon'][key], dec):<16}" for r in resultados))

    for r in resultados:
        print(f"\n--- FASE 2: Perturbaciones en {r['planta']} "
              f"(umbral de deteccion = {r['umbral']:.2f} C) ---")
        if not r["perturbaciones"]:
            print("No se detectaron perturbaciones.")
        for p in r["perturbaciones"]:
            print(f"\n  #{p['n']}  {TXT_TIPO[p['tipo']]}   "
                  f"ventana t = {p['t_inicio']:.0f} -> {p['t_fin']:.0f} s")
            for label, key, dec in FILAS_PERT[4:]:
                print(f"     {label:<26} {_fmt_celda(p[key], dec)}")
        for a in r["avisos"]:
            print(f"  [!] {a}")


# ----------------------------------------------------------------------------
# 7. GRAFICAS
# ----------------------------------------------------------------------------
def _sombrear(ax, t, fases, etiquetar=False):
    for f in fases:
        ax.axvspan(t[f["i0"]], t[f["i1"] - 1], color=COLOR_FASE[f["tipo"]], alpha=0.12, lw=0)
        if etiquetar:
            ax.text(t[f["i0"]], 1.01, f["nombre"], transform=ax.get_xaxis_transform(),
                    fontsize=8, color=COLOR_FASE[f["tipo"]], va="bottom")


def crear_figura_control(t, plantas):
    """plantas: lista de dict(res=..., pv=..., sp=..., pwm=..., corr=...). Una columna por planta."""
    n = len(plantas)
    fig, axs = plt.subplots(4, n, figsize=(7.5 * n, 11), squeeze=False, sharex=True)
    for c, P in enumerate(plantas):
        r = P["res"]
        a = axs[0, c]
        a.plot(t, P["sp"], "k--", alpha=0.7, label="SetPoint")
        a.plot(t, P["pv"], color="#222222", lw=1.6, label="Temperatura")
        _sombrear(a, t, r["fases"], etiquetar=True)
        for p in r["perturbaciones"]:
            mk = "v" if p["tipo"] == "BAJA" else "^"
            a.plot(p["t_ext"], p["y_ext"], mk, color=COLOR_FASE[p["tipo"]], ms=9)
            a.annotate(f"{p['desviacion']:+.1f} C", (p["t_ext"], p["y_ext"]),
                       textcoords="offset points", xytext=(6, -12 if p["tipo"] == "BAJA" else 8),
                       fontsize=8, color=COLOR_FASE[p["tipo"]])
        e = r["escalon"]
        a.set_title(f"{r['planta']}: respuesta por fases (OS={e['sobreimpulso_pct']:.1f}%)", pad=18)
        a.set_ylabel("Temperatura [C]")
        a.grid(True, ls=":", alpha=0.6)
        a.legend(loc="lower right", fontsize=8)

        a = axs[1, c]
        a.plot(t, P["pwm"], "#555555", lw=1.4)
        _sombrear(a, t, r["fases"])
        a.axhline(100, color="gray", ls=":")
        a.set_ylim(-5, 105)
        a.set_ylabel("PWM [%]")
        a.set_title("Accion de control")
        a.grid(True, ls=":", alpha=0.6)

        a = axs[2, c]
        err = r["sp_obj"] - P["pv"]          # contra el SP objetivo, solo desde el escalon
        err[:r["idx_step"]] = np.nan
        a.plot(t, err, "#555555", lw=1.4)
        _sombrear(a, t, r["fases"])
        a.axhline(0, color="k", ls="--", lw=0.8)
        a.set_ylabel("Error [C]")
        a.set_title("e(t) = SP objetivo - PV")
        a.grid(True, ls=":", alpha=0.6)

        a = axs[3, c]
        if P["corr"] is not None and not np.all(np.isnan(P["corr"])):
            a.plot(t, P["corr"], "#555555", lw=1.4)
        _sombrear(a, t, r["fases"])
        a.set_ylabel("Corriente [mA]")
        a.set_xlabel("Tiempo [s]")
        a.set_title("Consumo de corriente")
        a.grid(True, ls=":", alpha=0.6)

    fig.legend(handles=[Patch(color=COLOR_FASE["ESCALON"], alpha=0.4, label="Escalon"),
                        Patch(color=COLOR_FASE["BAJA"], alpha=0.4, label="Perturbacion BAJA"),
                        Patch(color=COLOR_FASE["SUBE"], alpha=0.4, label="Perturbacion SUBE")],
               loc="upper center", ncol=3, frameon=False)
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    return fig


def crear_figura_segmentos(t, plantas):
    """Una fila por planta y una columna por fase (zoom a cada segmento)."""
    ncol = max(len(P["res"]["fases"]) for P in plantas)
    fig, axs = plt.subplots(len(plantas), ncol, figsize=(5.2 * ncol, 3.8 * len(plantas)),
                            squeeze=False)
    for r_i, P in enumerate(plantas):
        r = P["res"]
        for c in range(ncol):
            ax = axs[r_i, c]
            if c >= len(r["fases"]):
                ax.axis("off")
                continue
            f = r["fases"][c]
            sl = slice(f["i0"], f["i1"])
            col = COLOR_FASE[f["tipo"]]
            ax.plot(t[sl], P["pv"][sl], color=col, lw=1.8, label="Temperatura")
            ax.plot(t[sl], P["sp"][sl], "k--", lw=0.8, alpha=0.6)
            ax.set_ylabel("Temp [C]", color=col)
            ax.grid(True, ls=":", alpha=0.6)
            ax2 = ax.twinx()
            ax2.plot(t[sl], P["pwm"][sl], color="#888888", lw=1.1)
            ax2.set_ylabel("PWM [%]", color="#888888")
            ax2.set_ylim(-5, 105)
            if f["tipo"] == "ESCALON":
                e = r["escalon"]
                sub = f"OS={e['sobreimpulso_pct']:.1f}%  tr={format_val(e['t_rise'], '.0f', 'n/d')} s"
            else:
                p = r["perturbaciones"][c - 1]
                sub = f"dT={p['desviacion']:+.1f}C  dPWM={p['delta_pwm']:+.1f}%"
            ax.set_title(f"{r['planta']} - {f['nombre']}\n{sub}", fontsize=10)
            ax.set_xlabel("Tiempo [s]")
    fig.tight_layout()
    return fig


def fig_a_buffer(fig, ancho_px=900, dpi=100):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    buf.seek(0)
    alto_px = int(ancho_px * fig.get_figheight() / fig.get_figwidth())
    return buf, ancho_px, alto_px


# ----------------------------------------------------------------------------
# 8. EXPORTACION A EXCEL
# ----------------------------------------------------------------------------
_HDR_FILL = PatternFill("solid", fgColor="DDEBF7")


def _cabecera(ws, fila, textos, col0=1):
    for j, tx in enumerate(textos):
        c = ws.cell(row=fila, column=col0 + j, value=tx)
        c.font = Font(bold=True)
        c.fill = _HDR_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def exportar_analisis_excel(resultados, plantas, t, df_params=None, buffers=None,
                            archivo_analizar="Control.xlsx",
                            carpeta_salida=os.path.join("AnalisisDatos", "Control")):
    os.makedirs(carpeta_salida, exist_ok=True)
    base = os.path.splitext(os.path.basename(archivo_analizar))[0]
    ruta = os.path.join(carpeta_salida, f"AnalisisControl_{base}.xlsx")
    wb = openpyxl.Workbook()

    # Hoja 1: escalon
    ws = wb.active
    ws.title = "Escalon"
    _cabecera(ws, 1, ["Metrica"] + [r["planta"] for r in resultados])
    for i, (label, key, dec) in enumerate(FILAS_ESCALON, start=2):
        ws.cell(row=i, column=1, value=label)
        for j, r in enumerate(resultados, start=2):
            v = limpiar_numero(r["escalon"][key], dec)
            if v is not None:
                ws.cell(row=i, column=j, value=v)
    ws.column_dimensions["A"].width = 34
    for j in range(len(resultados)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(2 + j)].width = 18

    # Hoja 2: perturbaciones (una fila por evento)
    wp = wb.create_sheet("Perturbaciones")
    _cabecera(wp, 1, ["Planta"] + [f[0] for f in FILAS_PERT])
    fila = 2
    for r in resultados:
        for p in r["perturbaciones"]:
            wp.cell(row=fila, column=1, value=r["planta"])
            for j, (_, key, dec) in enumerate(FILAS_PERT, start=2):
                v = p[key]
                v = TXT_TIPO.get(v, v) if isinstance(v, str) else limpiar_numero(v, dec)
                if v is not None:
                    wp.cell(row=fila, column=j, value=v)
            fila += 1
    if fila == 2:
        wp.cell(row=2, column=1, value="No se detectaron perturbaciones")
    for j in range(1, len(FILAS_PERT) + 2):
        wp.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 16

    # Hoja 3: parametros PID
    if df_params is not None and not df_params.empty:
        wpid = wb.create_sheet("PID")
        _cabecera(wpid, 1, list(df_params.columns))
        for i, (_, row) in enumerate(df_params.iterrows(), start=2):
            for j, h in enumerate(df_params.columns, start=1):
                v = limpiar_numero(row[h], 4)
                if v is not None:
                    wpid.cell(row=i, column=j, value=v)

    # Hoja 4: datos segmentados (remuestreados + fase de cada planta)
    wd = wb.create_sheet("Segmentos")
    cab = ["Tiempo"]
    for P in plantas:
        n = P["res"]["planta"]
        cab += [f"{n} Temp", f"{n} SP", f"{n} PWM", f"{n} Corriente", f"{n} Fase"]
    _cabecera(wd, 1, cab)
    for i in range(len(t)):
        fila_d = [float(t[i])]
        for P in plantas:
            c = None if P["corr"] is None or np.isnan(P["corr"][i]) else float(P["corr"][i])
            fila_d += [float(P["pv"][i]), float(P["sp"][i]), float(P["pwm"][i]), c,
                       str(P["res"]["etiqueta"][i])]
        wd.append(fila_d)
    for j in range(1, len(cab) + 1):
        wd.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 15

    # Hoja 5: graficas
    if buffers:
        wg = wb.create_sheet("Graficas")
        fila_img = 1
        for buf, w, h in buffers:
            buf.seek(0)
            img = XLImage(buf)
            img.width, img.height = w, h
            wg.add_image(img, f"A{fila_img}")
            fila_img += int(h / 20) + 3

    wb.save(ruta)
    return ruta


# ----------------------------------------------------------------------------
# 9. PROGRAMA PRINCIPAL
# ----------------------------------------------------------------------------
def main(path_documento=None):
    if path_documento is None:
        path_documento = sys.argv[1] if len(sys.argv) > 1 else seleccionar_archivo_datos("dataControl")
    if not path_documento:
        return
    print(f"\nProcesando archivo: {os.path.basename(path_documento)}")

    df_raw = pd.read_excel(path_documento)
    df = preparar_datos(df_raw, CFG["dt"])
    t = df["Tiempo"].to_numpy(dtype=float)
    print(f"Registro: {len(df_raw)} filas originales -> {len(df)} muestras uniformes "
          f"(dt = {CFG['dt']:.1f} s, {t[-1] - t[0]:.0f} s)")

    def col(nombre):
        return df[nombre].to_numpy(dtype=float) if nombre in df.columns else None

    plantas = []
    sp1 = "SetPoint1" if "SetPoint1" in df.columns else "SetPoint"
    plantas.append(dict(nombre="Q1", pv=col("Temperatura1"), sp=col(sp1), pwm=col("PWM1"),
                        corr=col("Corriente1")))
    if all(c in df.columns for c in ("Temperatura2", "PWM2", "SetPoint2")):
        plantas.append(dict(nombre="Q2", pv=col("Temperatura2"), sp=col("SetPoint2"),
                            pwm=col("PWM2"), corr=col("Corriente2")))

    for P in plantas:
        P["res"] = analizar_planta(t, P["pv"], P["sp"], P["pwm"], P["corr"], P["nombre"], CFG)
    resultados = [P["res"] for P in plantas]

    cols_pid = [c for c in ["Planta", "Kp", "Ki", "Kd", "Ts", "Cambio (s)"] if c in df_raw.columns]
    df_params = df_raw[cols_pid].dropna(how="all").copy() if cols_pid else None
    if df_params is not None and not df_params.empty and "Planta" not in df_params.columns:
        df_params.insert(0, "Planta", ["Q1"] * len(df_params))

    imprimir_resumen_control(resultados, df_params)

    print("\nGenerando graficas...")
    fig1 = crear_figura_control(t, plantas)
    fig2 = crear_figura_segmentos(t, plantas)
    buffers = [fig_a_buffer(fig1), fig_a_buffer(fig2)]

    ruta = exportar_analisis_excel(resultados, plantas, t, df_params, buffers,
                                   archivo_analizar=os.path.basename(path_documento))
    print(f"Reporte Excel generado en: {ruta}")
    print("Mostrando graficas interactivas...")
    plt.show()


if __name__ == "__main__":
    main()