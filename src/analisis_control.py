import os
import glob
import io
from datetime import datetime
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image
import matplotlib.pyplot as plt

def limpiar_consola():
    os.system('cls' if os.name == 'nt' else 'clear')

def seleccionar_archivo_datos(carpeta_datos="dataControl"):
    limpiar_consola()
    archivos = glob.glob(os.path.join(carpeta_datos, "*.xlsx"))
    if not archivos:
        print(f"No se encontraron archivos Excel en la carpeta '{carpeta_datos}'.")
        return None
    
    print("Archivos disponibles en dataControl:")
    for i, arch in enumerate(archivos):
        print(f"[{i}] {os.path.basename(arch)}")
    
    if len(archivos) == 1:
        print(f"Seleccionando unico archivo: {os.path.basename(archivos[0])}")
        return archivos[0]

    while True:
        try:
            seleccion = int(input(f"Seleccione el numero del archivo (0-{len(archivos)-1}): "))
            if 0 <= seleccion < len(archivos):
                return archivos[seleccion]
            print("Seleccion invalida.")
        except ValueError:
            print("Por favor ingrese un numero valido.")

def detectar_y_analizar_perturbacion(tiempo, pv, sp, pwm, idx_step):
    t_post = tiempo[idx_step:]
    pv_post = pv[idx_step:]
    pwm_post = pwm[idx_step:]
    sp_post = sp[idx_step:]

    if len(t_post) < 150:
        return None

    # Filtrar tiempo despues del transitorio inicial
    mask_estable = (t_post - t_post[0] > 180)
    if not np.any(mask_estable):
        return None

    # Detectar caida de temperatura bajo SP (> 1.0 °C)
    indices_caida = np.where(mask_estable & (pv_post <= (sp_post[0] - 1.0)))[0]
    if len(indices_caida) == 0:
        return None

    idx_inicio = indices_caida[0]
    t_inicio = t_post[idx_inicio]

    # Estado previo a la perturbacion
    idx_pre = max(0, idx_inicio - 50)
    y_pre = np.mean(pv_post[idx_pre:idx_inicio])
    pwm_pre = np.mean(pwm_post[idx_pre:idx_inicio])

    # Segmento durante la perturbacion
    t_pert = t_post[idx_inicio:]
    pv_pert = pv_post[idx_inicio:]
    pwm_pert = pwm_post[idx_inicio:]

    y_min = np.min(pv_pert)
    idx_min = np.argmin(pv_pert)
    t_min = t_pert[idx_min]
    caida_temp = y_pre - y_min

    pwm_max = np.max(pwm_pert)
    idx_pwm_max = np.argmax(pwm_pert)
    t_pwm_max = t_pert[idx_pwm_max]
    delta_pwm = pwm_max - pwm_pre

    # Tiempo de recuperacion (a 0.5 °C del SP)
    banda_recov = sp_post[0] - 0.5
    indices_recov = np.where((t_pert > t_min) & (pv_pert >= banda_recov))[0]
    t_recov = t_pert[indices_recov[0]] if len(indices_recov) > 0 else np.nan
    t_recup = (t_recov - t_inicio) if not np.isnan(t_recov) else np.nan

    # Estado final de la perturbacion
    n_fin = min(30, len(pv_pert))
    y_final = np.mean(pv_pert[-n_fin:])
    pwm_final = np.mean(pwm_pert[-n_fin:])

    # Indices de error en perturbacion
    dt_pert = np.diff(t_pert, prepend=t_pert[0])
    dt_pert[0] = dt_pert[1] if len(dt_pert) > 1 else 1.0
    err_pert = sp_post[0] - pv_pert
    iae_pert = np.sum(np.abs(err_pert) * dt_pert)
    ise_pert = np.sum((err_pert ** 2) * dt_pert)

    return {
        "detectada": True,
        "t_inicio": t_inicio,
        "y_pre": y_pre,
        "pwm_pre": pwm_pre,
        "y_min": y_min,
        "t_min": t_min,
        "caida_temp": caida_temp,
        "pwm_max": pwm_max,
        "t_pwm_max": t_pwm_max,
        "delta_pwm": delta_pwm,
        "t_recov": t_recov,
        "tiempo_recuperacion": t_recup,
        "y_final": y_final,
        "pwm_final": pwm_final,
        "iae_pert": iae_pert,
        "ise_pert": ise_pert
    }

def calcular_indices_y_metricas(tiempo, pv, sp, pwm, corriente=None, nombre_planta="Planta"):
    # Deteccion del escalon de SetPoint
    indices_cambio = np.where(sp > 0)[0]
    if len(indices_cambio) == 0:
        diff_sp = np.diff(sp)
        indices_cambio = np.where(diff_sp != 0)[0]
        idx_step = indices_cambio[0] + 1 if len(indices_cambio) > 0 else 0
    else:
        idx_step = indices_cambio[0]

    t_step = tiempo[idx_step]
    sp_final = sp[idx_step]
    y_inicial = pv[idx_step - 1] if idx_step > 0 else pv[0]
    delta_sp = sp_final - y_inicial

    # Segmento tras el escalon
    t_seg = tiempo[idx_step:] - t_step
    pv_seg = pv[idx_step:]
    sp_seg = sp[idx_step:]
    pwm_seg = pwm[idx_step:]
    error_seg = sp_seg - pv_seg

    # Metricas de respuesta temporal
    y_max = np.max(pv_seg)
    idx_max = np.argmax(pv_seg)
    t_peak = t_seg[idx_max]
    sobreimpulso_abs = max(0.0, y_max - sp_final)
    sobreimpulso_pct = (sobreimpulso_abs / delta_sp) * 100.0 if delta_sp != 0 else 0.0

    # Tiempo de subida tr (10% a 90%)
    val_10 = y_inicial + 0.10 * delta_sp
    val_90 = y_inicial + 0.90 * delta_sp
    idx_10 = np.where(pv_seg >= val_10)[0]
    idx_90 = np.where(pv_seg >= val_90)[0]
    t_rise = (t_seg[idx_90[0]] - t_seg[idx_10[0]]) if (len(idx_10) > 0 and len(idx_90) > 0) else np.nan

    # Tiempo de establecimiento ts (2% y 5%)
    banda_2 = 0.02 * delta_sp
    banda_5 = 0.05 * delta_sp
    f_banda_2 = np.where(np.abs(error_seg) > banda_2)[0]
    f_banda_5 = np.where(np.abs(error_seg) > banda_5)[0]

    t_settling_2 = t_seg[f_banda_2[-1] + 1] if len(f_banda_2) > 0 and (f_banda_2[-1] + 1) < len(t_seg) else (t_seg[-1] if len(f_banda_2) == 0 else np.nan)
    t_settling_5 = t_seg[f_banda_5[-1] + 1] if len(f_banda_5) > 0 and (f_banda_5[-1] + 1) < len(t_seg) else (t_seg[-1] if len(f_banda_5) == 0 else np.nan)

    # Estado estacionario y error
    n_ss = min(10, len(pv_seg))
    y_ss = np.mean(pv_seg[-n_ss:])
    error_ss = sp_final - y_ss

    # Indices de desempeno integrales
    dt = np.diff(t_seg, prepend=t_seg[0])
    dt[0] = dt[1] if len(dt) > 1 else 1.0
    iae = np.sum(np.abs(error_seg) * dt)
    ise = np.sum((error_seg ** 2) * dt)
    itae = np.sum(t_seg * np.abs(error_seg) * dt)

    # Esfuerzo de control y suavidad
    pwm_max = np.max(pwm_seg)
    pwm_ss = np.mean(pwm_seg[-n_ss:])
    tv_u = np.sum(np.abs(np.diff(pwm_seg)))

    # Corriente
    i_max = np.nan
    i_ss = np.nan
    if corriente is not None:
        corr_seg = corriente[idx_step:]
        i_max = np.max(corr_seg)
        i_ss = np.mean(corr_seg[-n_ss:])

    # Deteccion de perturbaciones
    info_pert = detectar_y_analizar_perturbacion(tiempo, pv, sp, pwm, idx_step)

    return {
        "planta": nombre_planta,
        "idx_step": idx_step,
        "t_step": t_step,
        "y_inicial": y_inicial,
        "sp_final": sp_final,
        "delta_sp": delta_sp,
        "y_max": y_max,
        "t_peak": t_peak,
        "sobreimpulso_abs": sobreimpulso_abs,
        "sobreimpulso_pct": sobreimpulso_pct,
        "t_rise": t_rise,
        "t_settling_2": t_settling_2,
        "t_settling_5": t_settling_5,
        "y_ss": y_ss,
        "error_ss": error_ss,
        "iae": iae,
        "ise": ise,
        "itae": itae,
        "pwm_max": pwm_max,
        "pwm_ss": pwm_ss,
        "tv_u": tv_u,
        "i_max": i_max,
        "i_ss": i_ss,
        "perturbacion": info_pert
    }

def format_val(val, fmt=".2f", default="-"):
    if val is None or np.isnan(val):
        return default
    return f"{val:{fmt}}"

def limpiar_numero(x, dec=2):
    if x is None:
        return None
    try:
        if np.isnan(x):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(x, (int, np.integer)):
        return int(x)
    if isinstance(x, (float, np.floating)):
        r = round(float(x), dec)
        return int(r) if r.is_integer() else r
    return x

def imprimir_resumen_control(m1, m2=None, df_params=None):
    print("\n" + "=" * 55)
    print("          RESUMEN DEL CONTROL")
    print("=" * 55)

    if df_params is not None and not df_params.empty:
        print("\n--- Parametros PID (Configurados) ---")
        print(df_params.to_string(index=False))

    print("\n--- Metricas de Desempeno ---")
    tiene_m2 = m2 is not None
    if tiene_m2:
        hdr = f"{'Metrica':<28} | {'Q1':<10} | {'Q2':<10}"
        sep = "-" * len(hdr)
        print(hdr)
        print(sep)
        rows = [
            ("SetPoint [°C]", format_val(m1['sp_final']), format_val(m2['sp_final'])),
            ("Sobreimpulso [%]", format_val(m1['sobreimpulso_pct']), format_val(m2['sobreimpulso_pct'])),
            ("Tiempo Subida tr [s]", format_val(m1['t_rise']), format_val(m2['t_rise'])),
            ("Tiempo Establ. 5% [s]", format_val(m1['t_settling_5']), format_val(m2['t_settling_5'])),
            ("Tiempo Establ. 2% [s]", format_val(m1['t_settling_2']), format_val(m2['t_settling_2'])),
            ("Error Estado Estable [°C]", format_val(m1['error_ss']), format_val(m2['error_ss'])),
            ("IAE", format_val(m1['iae']), format_val(m2['iae'])),
            ("ISE", format_val(m1['ise']), format_val(m2['ise'])),
            ("ITAE", format_val(m1['itae']), format_val(m2['itae'])),
            ("PWM Max [%]", format_val(m1['pwm_max']), format_val(m2['pwm_max'])),
            ("PWM Estable [%]", format_val(m1['pwm_ss']), format_val(m2['pwm_ss'])),
            ("Suavidad TV(u)", format_val(m1['tv_u']), format_val(m2['tv_u'])),
            ("Corriente Max [mA]", format_val(m1['i_max'], ".1f"), format_val(m2['i_max'], ".1f")),
        ]
        for name, v1, v2 in rows:
            print(f"{name:<28} | {v1:<10} | {v2:<10}")
        print(sep)
    else:
        hdr = f"{'Metrica':<28} | {'Q1':<10}"
        sep = "-" * len(hdr)
        print(hdr)
        print(sep)
        rows = [
            ("SetPoint [°C]", format_val(m1['sp_final'])),
            ("Sobreimpulso [%]", format_val(m1['sobreimpulso_pct'])),
            ("Tiempo Subida tr [s]", format_val(m1['t_rise'])),
            ("Tiempo Establ. 5% [s]", format_val(m1['t_settling_5'])),
            ("Tiempo Establ. 2% [s]", format_val(m1['t_settling_2'])),
            ("Error Estado Estable [°C]", format_val(m1['error_ss'])),
            ("IAE", format_val(m1['iae'])),
            ("ISE", format_val(m1['ise'])),
            ("ITAE", format_val(m1['itae'])),
            ("PWM Max [%]", format_val(m1['pwm_max'])),
            ("PWM Estable [%]", format_val(m1['pwm_ss'])),
            ("Suavidad TV(u)", format_val(m1['tv_u'])),
            ("Corriente Max [mA]", format_val(m1['i_max'], ".1f")),
        ]
        for name, v1 in rows:
            print(f"{name:<28} | {v1:<10}")
        print(sep)

    # Analisis de perturbacion externa si existe
    p1 = m1.get('perturbacion')
    p2 = m2.get('perturbacion') if tiene_m2 else None
    if p1 or p2:
        print("\n--- Rechazo a Perturbaciones ---")
        t_inicio = p1['t_inicio'] if p1 else p2['t_inicio']
        print(f"Inicio deteccion: t ≈ {t_inicio:.1f} s")
        if tiene_m2 and p1 and p2:
            hdr_p = f"{'Metrica':<28} | {'Q1':<10} | {'Q2':<10}"
            print(hdr_p)
            print("-" * len(hdr_p))
            print(f"{'Caida Temp ΔT [°C]':<28} | {p1['caida_temp']:<10.2f} | {p2['caida_temp']:<10.2f}")
            print(f"{'PWM Max Reaccion [%]':<28} | {p1['pwm_max']:<10.2f} | {p2['pwm_max']:<10.2f}")
            print(f"{'Tiempo Recuperacion [s]':<28} | {format_val(p1['tiempo_recuperacion']):<10} | {format_val(p2['tiempo_recuperacion']):<10}")
            print(f"{'IAE Perturbacion':<28} | {p1['iae_pert']:<10.2f} | {p2['iae_pert']:<10.2f}")
            print("-" * len(hdr_p))
        elif p1:
            hdr_p = f"{'Metrica':<28} | {'Q1':<10}"
            print(hdr_p)
            print("-" * len(hdr_p))
            print(f"{'Caida Temp ΔT [°C]':<28} | {p1['caida_temp']:<10.2f}")
            print(f"{'PWM Max Reaccion [%]':<28} | {p1['pwm_max']:<10.2f}")
            print(f"{'Tiempo Recuperacion [s]':<28} | {format_val(p1['tiempo_recuperacion']):<10}")
            print(f"{'IAE Perturbacion':<28} | {p1['iae_pert']:<10.2f}")
            print("-" * len(hdr_p))

def exportar_analisis_excel(m1, m2=None, df_params=None, img_buf=None, carpeta_salida=os.path.join("AnalisisDatos", "Control")):
    os.makedirs(carpeta_salida, exist_ok=True)
    fecha_actual = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre_archivo = f"AnalisisControl_{fecha_actual}.xlsx"
    ruta_archivo = os.path.join(carpeta_salida, nombre_archivo)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AnalisisControl"

    # Definicion de filas de metricas
    filas_metricas = [
        ("planta", m1["planta"], m2["planta"] if m2 else None, 0),
        ("Setpoint", m1["sp_final"], m2["sp_final"] if m2 else None, 2),
        ("Delta SP [°C]", m1["delta_sp"], m2["delta_sp"] if m2 else None, 2),
        ("Temp Max [°C]", m1["y_max"], m2["y_max"] if m2 else None, 2),
        ("Tiempo Pico [s]", m1["t_peak"], m2["t_peak"] if m2 else None, 1),
        ("Sobreimpulso [°C]", m1["sobreimpulso_abs"], m2["sobreimpulso_abs"] if m2 else None, 2),
        ("Sobreimpulso [%]", m1["sobreimpulso_pct"], m2["sobreimpulso_pct"] if m2 else None, 2),
        ("Tiempo Subida tr [s]", m1["t_rise"], m2["t_rise"] if m2 else None, 2),
        ("Tiempo Establ 5% [s]", m1["t_settling_5"], m2["t_settling_5"] if m2 else None, 2),
        ("Tiempo Establ 2% [s]", m1["t_settling_2"], m2["t_settling_2"] if m2 else None, 2),
        ("Temp Final Yss [°C]", m1["y_ss"], m2["y_ss"] if m2 else None, 2),
        ("Error Regimen ess [°C]", m1["error_ss"], m2["error_ss"] if m2 else None, 2),
        ("IAE", m1["iae"], m2["iae"] if m2 else None, 2),
        ("ISE", m1["ise"], m2["ise"] if m2 else None, 3),
        ("ITAE", m1["itae"], m2["itae"] if m2 else None, 2),
        ("PWM Maximo [%]", m1["pwm_max"], m2["pwm_max"] if m2 else None, 2),
        ("PWM Estable [%]", m1["pwm_ss"], m2["pwm_ss"] if m2 else None, 2),
        ("Variacion Total TV(u)", m1["tv_u"], m2["tv_u"] if m2 else None, 2),
        ("Corriente Max [mA]", m1["i_max"], m2["i_max"] if m2 else None, 1),
        ("Corriente Estable [mA]", m1["i_ss"], m2["i_ss"] if m2 else None, 1),
    ]

    for r_idx, (label, val_q1, val_q2, dec) in enumerate(filas_metricas, start=1):
        ws.cell(row=r_idx, column=1, value=label)
        v1 = limpiar_numero(val_q1, dec)
        if v1 is not None:
            ws.cell(row=r_idx, column=2, value=v1)
        if m2 is not None:
            v2 = limpiar_numero(val_q2, dec)
            if v2 is not None:
                ws.cell(row=r_idx, column=3, value=v2)

    # Parametros PID en columnas E:J
    if df_params is not None and not df_params.empty:
        cols_pid = [c for c in ['Planta', 'Kp', 'Ki', 'Kd', 'Ts', 'Cambio (s)'] if c in df_params.columns]
        for c_idx, h in enumerate(cols_pid, start=5):
            ws.cell(row=1, column=c_idx, value=h)

        for f_idx, (_, row_data) in enumerate(df_params.iterrows(), start=2):
            for c_idx, h in enumerate(cols_pid, start=5):
                val = limpiar_numero(row_data[h], dec=4)
                if val is not None:
                    ws.cell(row=f_idx, column=c_idx, value=val)

    # Metricas de perturbacion si existen
    p1 = m1.get("perturbacion")
    p2 = m2.get("perturbacion") if m2 else None
    if p1 or p2:
        start_pert = len(filas_metricas) + 2
        ws.cell(row=start_pert, column=1, value="--- Perturbacion ---")
        filas_p = [
            ("Tiempo Inicio [s]", p1["t_inicio"] if p1 else None, p2["t_inicio"] if p2 else None, 1),
            ("Temp Previa [°C]", p1["y_pre"] if p1 else None, p2["y_pre"] if p2 else None, 2),
            ("Temp Min Nadir [°C]", p1["y_min"] if p1 else None, p2["y_min"] if p2 else None, 2),
            ("Caida Temp ΔT [°C]", p1["caida_temp"] if p1 else None, p2["caida_temp"] if p2 else None, 2),
            ("PWM Max Reaccion [%]", p1["pwm_max"] if p1 else None, p2["pwm_max"] if p2 else None, 2),
            ("Delta PWM [%]", p1["delta_pwm"] if p1 else None, p2["delta_pwm"] if p2 else None, 2),
            ("Tiempo Recuperacion [s]", p1["tiempo_recuperacion"] if p1 else None, p2["tiempo_recuperacion"] if p2 else None, 2),
            ("Temp Final Pert [°C]", p1["y_final"] if p1 else None, p2["y_final"] if p2 else None, 2),
            ("IAE Perturbacion", p1["iae_pert"] if p1 else None, p2["iae_pert"] if p2 else None, 2),
            ("ISE Perturbacion", p1["ise_pert"] if p1 else None, p2["ise_pert"] if p2 else None, 2),
        ]
        for offset, (label, val_p1, val_p2, dec) in enumerate(filas_p, start=1):
            curr_row = start_pert + offset
            ws.cell(row=curr_row, column=1, value=label)
            vp1 = limpiar_numero(val_p1, dec)
            if vp1 is not None:
                ws.cell(row=curr_row, column=2, value=vp1)
            if m2 is not None:
                vp2 = limpiar_numero(val_p2, dec)
                if vp2 is not None:
                    ws.cell(row=curr_row, column=3, value=vp2)

    # Anchos de columna
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 14
    if m2 is not None:
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 4
    else:
        ws.column_dimensions["C"].width = 4
        ws.column_dimensions["D"].width = 4
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 10
    ws.column_dimensions["J"].width = 12

    # Incrustar imagen de la grafica
    if img_buf is not None:
        img_buf.seek(0)
        img = Image(img_buf)
        img.width = 720
        img.height = 480
        ws.add_image(img, "E5")

    wb.save(ruta_archivo)
    return ruta_archivo

def crear_figura_control(tiempo, pv1, sp1, pwm1, corr1,
                         pv2=None, sp2=None, pwm2=None, corr2=None,
                         m1=None, m2=None):
    fig, axs = plt.subplots(2, 2, figsize=(12, 8))
    
    # Tiempo de perturbacion si existe
    p1 = m1.get('perturbacion') if m1 else None
    p2 = m2.get('perturbacion') if m2 else None
    t_pert = p1['t_inicio'] if p1 else (p2['t_inicio'] if p2 else None)

    # [0, 0] Respuesta Temporal
    lbl_sp1 = 'SetPoint Q1' if pv2 is not None else 'SetPoint'
    axs[0, 0].plot(tiempo, sp1, 'k--', label=lbl_sp1, alpha=0.7)
    axs[0, 0].plot(tiempo, pv1, 'r-', linewidth=1.8, label=f'Temp Q1 (OS={m1["sobreimpulso_pct"]:.1f}%)')
    if pv2 is not None:
        if not np.array_equal(sp1, sp2):
            axs[0, 0].plot(tiempo, sp2, 'c--', label='SetPoint Q2', alpha=0.6)
        axs[0, 0].plot(tiempo, pv2, 'b-', linewidth=1.8, label=f'Temp Q2 (OS={m2["sobreimpulso_pct"]:.1f}%)')

    if t_pert is not None:
        axs[0, 0].axvline(t_pert, color='purple', linestyle=':', linewidth=1.5, label=f'Perturbacion (~{t_pert:.0f}s)')

    axs[0, 0].set_title('Respuesta Temporal del Sistema (Lazo Cerrado)')
    axs[0, 0].set_ylabel('Temperatura [°C]')
    axs[0, 0].legend()
    axs[0, 0].grid(True, linestyle=':', alpha=0.6)

    # [1, 0] Esfuerzo de Control (PWM)
    axs[1, 0].plot(tiempo, pwm1, 'r-', label='PWM Q1', linewidth=1.4)
    if pwm2 is not None:
        axs[1, 0].plot(tiempo, pwm2, 'b-', label='PWM Q2', linewidth=1.4)
    if t_pert is not None:
        axs[1, 0].axvline(t_pert, color='purple', linestyle=':', linewidth=1.5)
    axs[1, 0].axhline(100, color='gray', linestyle=':', label='Saturacion 100%')
    axs[1, 0].set_title('Accion de Control / Esfuerzo (PWM)')
    axs[1, 0].set_ylabel('Ciclo de Trabajo [%]')
    axs[1, 0].set_xlabel('Tiempo [s]')
    axs[1, 0].set_ylim(-5, 105)
    axs[1, 0].legend()
    axs[1, 0].grid(True, linestyle=':', alpha=0.6)

    # [0, 1] Error de Control
    err1 = sp1 - pv1
    axs[0, 1].plot(tiempo, err1, 'r-', label=f'Error Q1 (IAE={m1["iae"]:.1f})', linewidth=1.5)
    if pv2 is not None:
        err2 = sp2 - pv2
        axs[0, 1].plot(tiempo, err2, 'b-', label=f'Error Q2 (IAE={m2["iae"]:.1f})', linewidth=1.5)
    if t_pert is not None:
        axs[0, 1].axvline(t_pert, color='purple', linestyle=':', linewidth=1.5)
    axs[0, 1].axhline(0, color='black', linestyle='--', linewidth=0.8)
    axs[0, 1].set_title('Senal de Error e(t) = SP - PV')
    axs[0, 1].set_ylabel('Error [°C]')
    axs[0, 1].legend()
    axs[0, 1].grid(True, linestyle=':', alpha=0.6)

    # [1, 1] Consumo de Corriente
    if corr1 is not None and not np.all(np.isnan(corr1)):
        axs[1, 1].plot(tiempo, corr1, 'r-', label='Corriente Q1', linewidth=1.4)
    if corr2 is not None and not np.all(np.isnan(corr2)):
        axs[1, 1].plot(tiempo, corr2, 'b-', label='Corriente Q2', linewidth=1.4)
    if t_pert is not None:
        axs[1, 1].axvline(t_pert, color='purple', linestyle=':', linewidth=1.5)
    axs[1, 1].set_title('Consumo de Corriente Electrica')
    axs[1, 1].set_ylabel('Corriente [mA]')
    axs[1, 1].set_xlabel('Tiempo [s]')
    axs[1, 1].legend()
    axs[1, 1].grid(True, linestyle=':', alpha=0.6)

    plt.tight_layout()

    # Guardar imagen en memoria para incrustar en Excel
    img_buf = io.BytesIO()
    fig.savefig(img_buf, format='png', dpi=100, bbox_inches='tight')
    img_buf.seek(0)

    return fig, img_buf

def main():
    # 1. Seleccion de archivo
    path_documento = seleccionar_archivo_datos("dataControl")
    if not path_documento:
        return
    print(f"\nProcesando archivo: {os.path.basename(path_documento)}")

    # 2. Lectura y preparacion de datos
    df = pd.read_excel(path_documento)
    tiempo = np.double(df['Tiempo'])

    # Canales de Planta Q1
    pv1 = np.double(df['Temperatura1'])
    pwm1 = np.double(df['PWM1'])
    sp1_col = 'SetPoint1' if 'SetPoint1' in df.columns else 'SetPoint'
    sp1 = np.double(df[sp1_col])
    corr1 = np.double(df['Corriente1']) if 'Corriente1' in df.columns else None

    # Canales de Planta Q2 (si existen en el archivo)
    tiene_q2 = 'Temperatura2' in df.columns and 'PWM2' in df.columns and 'SetPoint2' in df.columns
    pv2 = np.double(df['Temperatura2']) if tiene_q2 else None
    pwm2 = np.double(df['PWM2']) if tiene_q2 else None
    sp2 = np.double(df['SetPoint2']) if tiene_q2 else None
    corr2 = np.double(df['Corriente2']) if (tiene_q2 and 'Corriente2' in df.columns) else None

    # Parametros PID del archivo
    cols_pid = [c for c in ['Planta', 'Kp', 'Ki', 'Kd', 'Ts', 'Cambio (s)'] if c in df.columns]
    df_params = df[cols_pid].dropna(how='all').copy() if cols_pid else None
    if df_params is not None and not df_params.empty and 'Planta' not in df_params.columns:
        df_params.insert(0, 'Planta', ['Q1'] * len(df_params))

    # 3. Calculo de metricas e indices
    m1 = calcular_indices_y_metricas(tiempo, pv1, sp1, pwm1, corr1, nombre_planta="Planta Q1")
    m2 = calcular_indices_y_metricas(tiempo, pv2, sp2, pwm2, corr2, nombre_planta="Planta Q2") if tiene_q2 else None

    # 4. Impresion de resultados
    imprimir_resumen_control(m1, m2, df_params)

    # 5. Generacion de graficas y buffer para Excel
    print("\nGenerando graficas del proceso de control...")
    fig, img_buf = crear_figura_control(tiempo, pv1, sp1, pwm1, corr1, pv2, sp2, pwm2, corr2, m1, m2)

    # 6. Exportacion a Excel con imagen incrustada
    ruta_excel = exportar_analisis_excel(m1, m2, df_params, img_buf=img_buf)
    print(f"Reporte Excel generado en: {ruta_excel}")

    # 7. Visualizacion interactiva
    print("Mostrando graficas interactivas...")
    plt.show()

if __name__ == "__main__":
    main()
